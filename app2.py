"""
Student Integrity Assessment Portal
-----------------------------------
Upload two student documents (.txt, .docx, .pdf) to:
1. Compare text similarity using word n-grams.
2. Show writing-style indicators that may resemble AI-generated text.

Important:
- Similarity is not proof of plagiarism.
- The AI-style score is a heuristic, NOT a reliable AI detector.
- Scanned/image-only PDFs may not contain extractable text.

Install:
    pip install streamlit python-docx PyPDF2

Run:
    streamlit run integrity_portal.py
"""

import re
import statistics
from typing import List, Set

import streamlit as st

# Optional file readers
try:
    import docx
except ImportError:
    docx = None

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None


# -------------------------------------------------------------------
# 1. FILE READING
# -------------------------------------------------------------------

def extract_text(uploaded_file) -> str:
    """Extract text from TXT, DOCX, or PDF uploads."""
    filename = uploaded_file.name.lower()

    if filename.endswith(".txt"):
        raw = uploaded_file.getvalue()
        # Try UTF-8 first, then a common Windows encoding.
        try:
            return raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            return raw.decode("cp1252", errors="replace")

    if filename.endswith(".docx"):
        if docx is None:
            raise RuntimeError(
                "python-docx is not installed. Run: pip install python-docx"
            )
        document = docx.Document(uploaded_file)
        paragraphs = [p.text for p in document.paragraphs]
        # Include table contents too.
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)

    if filename.endswith(".pdf"):
        if PdfReader is None:
            raise RuntimeError(
                "PyPDF2 is not installed. Run: pip install PyPDF2"
            )
        reader = PdfReader(uploaded_file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    raise ValueError("Unsupported file type. Use .txt, .docx, or .pdf.")


# -------------------------------------------------------------------
# 2. TEXT CLEANING
# -------------------------------------------------------------------

def clean_words(text: str) -> List[str]:
    """Normalize text into lowercase word tokens."""
    return re.findall(r"[a-z0-9]+(?:'[a-z0-9]+)?", text.lower())


def split_sentences(text: str) -> List[str]:
    """Split text approximately into sentences."""
    parts = re.split(r"(?<=[.!?])\s+|[\r\n]+", text.strip())
    return [part.strip() for part in parts if part.strip()]


# -------------------------------------------------------------------
# 3. TEXT SIMILARITY
# -------------------------------------------------------------------

def get_ngrams(words: List[str], n: int = 3) -> Set[str]:
    """Return unique word n-grams."""
    if n < 1:
        raise ValueError("n must be at least 1")
    if not words:
        return set()
    if len(words) < n:
        return {" ".join(words)}
    return {" ".join(words[i:i + n]) for i in range(len(words) - n + 1)}


def match_percentage(text_a: str, text_b: str, n: int = 3) -> float:
    """Jaccard similarity of unique word n-grams, expressed as a percentage."""
    ngrams_a = get_ngrams(clean_words(text_a), n)
    ngrams_b = get_ngrams(clean_words(text_b), n)

    if not ngrams_a or not ngrams_b:
        return 0.0

    union = ngrams_a | ngrams_b
    common = ngrams_a & ngrams_b
    return round(len(common) / len(union) * 100, 2)


def matching_sentences(
    text_a: str,
    text_b: str,
    n: int = 3,
    threshold: float = 40.0,
) -> List[str]:
    """Find sentences in A with substantial n-gram overlap against B."""
    ngrams_b = get_ngrams(clean_words(text_b), n)
    flagged = []

    if not ngrams_b:
        return flagged

    for sentence in split_sentences(text_a):
        words = clean_words(sentence)
        if len(words) < n:
            continue

        sentence_ngrams = get_ngrams(words, n)
        if not sentence_ngrams:
            continue

        overlap = len(sentence_ngrams & ngrams_b) / len(sentence_ngrams) * 100
        if overlap >= threshold:
            flagged.append(sentence)

    return flagged


# -------------------------------------------------------------------
# 4. WRITING-STYLE INDICATORS (HEURISTIC ONLY)
# -------------------------------------------------------------------

TRANSITION_PHRASES = [
    "furthermore",
    "additionally",
    "moreover",
    "in conclusion",
    "it is important to note",
    "overall",
    "in summary",
    "therefore",
    "as a result",
    "on the other hand",
    "in addition",
    "consequently",
]


def sentence_length_uniformity_score(text: str) -> float:
    """
    Heuristic score based on sentence-length consistency.
    This is not evidence that a text was written by AI.
    """
    lengths = [
        len(clean_words(sentence))
        for sentence in split_sentences(text)
        if clean_words(sentence)
    ]

    if len(lengths) < 3:
        return 50.0

    mean_length = statistics.mean(lengths)
    if mean_length == 0:
        return 50.0

    coefficient_of_variation = statistics.pstdev(lengths) / mean_length
    score = max(0.0, 100.0 - coefficient_of_variation * 100.0)
    return round(min(score, 100.0), 2)


def transition_phrase_score(text: str) -> float:
    """Heuristic based on listed transition phrases per 100 words."""
    words = clean_words(text)
    if not words:
        return 0.0

    lower_text = text.lower()
    count = sum(lower_text.count(phrase) for phrase in TRANSITION_PHRASES)
    rate = (count / len(words)) * 100 * 20
    return round(min(rate, 100.0), 2)


def vocabulary_diversity_score(text: str) -> float:
    """
    Convert unique-word ratio into a heuristic repetition score.
    This is not a validated AI-detection method.
    """
    words = clean_words(text)
    if not words:
        return 50.0

    unique_ratio = len(set(words)) / len(words) * 100

    if unique_ratio >= 70:
        return 10.0
    if unique_ratio <= 40:
        return 90.0

    return round(90.0 - ((unique_ratio - 40.0) / 30.0) * 80.0, 2)


def ai_likelihood_score(text: str) -> dict:
    """
    Return a heuristic writing-style indicator.
    Do not interpret this as the probability that text was AI-generated.
    """
    uniformity = sentence_length_uniformity_score(text)
    transitions = transition_phrase_score(text)
    vocabulary = vocabulary_diversity_score(text)

    overall = round(
        uniformity * 0.4 + transitions * 0.3 + vocabulary * 0.3,
        2,
    )

    return {
        "overall": overall,
        "sentence_uniformity": uniformity,
        "transition_phrases": transitions,
        "vocabulary_repetition": vocabulary,
    }


# -------------------------------------------------------------------
# 5. STREAMLIT UI
# -------------------------------------------------------------------

st.set_page_config(
    page_title="Student Integrity Assessment Portal",
    page_icon="🎓",
    layout="wide",
)

st.title("🎓 Student Integrity Assessment Portal")
st.caption(
    "Compare two student submissions for text similarity and inspect "
    "writing-style indicators. Scores are heuristic and require human review."
)

col_a, col_b = st.columns(2)

with col_a:
    file_a = st.file_uploader(
        "Upload Student A's file (.txt/.docx/.pdf)",
        type=["txt", "docx", "pdf"],
        key="file_a",
    )

with col_b:
    file_b = st.file_uploader(
        "Upload Student B's file (.txt/.docx/.pdf)",
        type=["txt", "docx", "pdf"],
        key="file_b",
    )

if st.button("Analyze Documents", type="primary"):
    if file_a is None or file_b is None:
        st.error("Please upload both files.")
    else:
        try:
            text_a = extract_text(file_a)
            text_b = extract_text(file_b)

            if not text_a.strip() or not text_b.strip():
                st.error(
                    "Could not extract readable text from one or both files. "
                    "For scanned PDFs, use a text-based PDF or OCR first."
                )
            else:
                match_pct = match_percentage(text_a, text_b)
                flagged = matching_sentences(text_a, text_b)
                ai_a = ai_likelihood_score(text_a)
                ai_b = ai_likelihood_score(text_b)

                st.divider()
                st.subheader("📊 Integrity Report")

                m1, m2, m3 = st.columns(3)
                m1.metric("Document Text Similarity", f"{match_pct}%")
                m2.metric("Student A – Style Indicator", f"{ai_a['overall']}%")
                m3.metric("Student B – Style Indicator", f"{ai_b['overall']}%")

                st.caption(
                    "The style indicator is not a probability of AI authorship "
                    "and must not be used alone to make academic-integrity decisions."
                )

                if match_pct >= 60:
                    st.error(
                        "High text similarity detected. Review the source text "
                        "and context manually."
                    )
                elif match_pct >= 30:
                    st.warning(
                        "Moderate text similarity detected. A manual review may help."
                    )
                else:
                    st.success(
                        "Low text similarity detected by this n-gram method."
                    )

                st.divider()
                col_x, col_y = st.columns(2)

                with col_x:
                    st.markdown("**Student A – Style Indicator Breakdown**")
                    st.write(
                        f"- Sentence-length uniformity: "
                        f"{ai_a['sentence_uniformity']}%"
                    )
                    st.write(
                        f"- Transition-phrase frequency: "
                        f"{ai_a['transition_phrases']}%"
                    )
                    st.write(
                        f"- Vocabulary-repetition indicator: "
                        f"{ai_a['vocabulary_repetition']}%"
                    )

                with col_y:
                    st.markdown("**Student B – Style Indicator Breakdown**")
                    st.write(
                        f"- Sentence-length uniformity: "
                        f"{ai_b['sentence_uniformity']}%"
                    )
                    st.write(
                        f"- Transition-phrase frequency: "
                        f"{ai_b['transition_phrases']}%"
                    )
                    st.write(
                        f"- Vocabulary-repetition indicator: "
                        f"{ai_b['vocabulary_repetition']}%"
                    )

                if flagged:
                    st.divider()
                    st.subheader("🔍 Potentially Matching Sentences")
                    st.caption(
                        "These sentences are from Student A and overlap with "
                        "Student B under the configured n-gram rule."
                    )
                    for sentence in flagged:
                        st.markdown(f"> {sentence}")
                else:
                    st.info(
                        "No sentences from Student A crossed the matching threshold."
                    )

                st.divider()
                st.info(
                    "**Important:** Similarity can occur in quotations, common "
                    "phrasing, templates, or properly cited material. These scores "
                    "are indicators, not proof of plagiarism or AI use."
                )

        except Exception as exc:
            st.error(f"Something went wrong: {exc}")
